from __future__ import annotations

import io
import os
from datetime import date, datetime

from starlette.concurrency import run_in_threadpool

from app.modules.documents.domain.metadata import MetadataExtractor

_PDF_EXTS = {".pdf"}
_DOCX_EXTS = {".docx"}
_XLSX_EXTS = {".xlsx", ".xlsm"}


def _ext(file_name: str | None) -> str:
    return os.path.splitext(file_name or "")[1].lower()


def _clean(value: object) -> object | None:
    """Coerce a property value into something JSON-serialisable, or drop it."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value).strip()
    return text or None


def _compact(props: dict) -> dict:
    """Drop keys whose value is None/empty after cleaning."""
    out: dict = {}
    for key, value in props.items():
        cleaned = _clean(value)
        if cleaned is not None:
            out[key] = cleaned
    return out


class CompositeMetadataExtractor(MetadataExtractor):
    """Derives content statistics + format-specific document properties.

    Pure-Python and CPU-light, but parsing the source bytes again (pypdf /
    python-docx / openpyxl) is synchronous, so the work is offloaded to a
    worker thread.
    """

    async def extract(
        self,
        *,
        data: bytes,
        text: str,
        file_name: str | None,
        mime_type: str | None,
    ) -> dict:
        return await run_in_threadpool(self._extract_sync, data, text, file_name, mime_type)

    def _extract_sync(
        self, data: bytes, text: str, file_name: str | None, mime_type: str | None
    ) -> dict:
        meta: dict = {"content": self._content_stats(text)}

        ext = _ext(file_name)
        mime = (mime_type or "").lower()
        try:
            if ext in _PDF_EXTS or mime == "application/pdf":
                props = self._pdf_props(data)
            elif ext in _DOCX_EXTS or "wordprocessingml" in mime:
                props = self._docx_props(data)
            elif ext in _XLSX_EXTS or "spreadsheetml" in mime:
                props = self._xlsx_props(data)
            else:
                props = {}
        except Exception as exc:  # never fail parsing because props are unreadable
            props = {"error": f"property extraction failed: {exc}"}

        if props:
            meta["properties"] = props
        return meta

    # --- content statistics --------------------------------------------

    def _content_stats(self, text: str) -> dict:
        stripped = text.strip()
        words = stripped.split()
        return {
            "char_count": len(text),
            "word_count": len(words),
            "line_count": text.count("\n") + 1 if stripped else 0,
            "language": self._detect_language(stripped),
        }

    def _detect_language(self, text: str) -> str | None:
        if len(text) < 20:  # too short to detect reliably
            return None
        try:
            from langdetect import DetectorFactory, detect

            DetectorFactory.seed = 0  # deterministic results
            return detect(text)
        except Exception:
            return None

    # --- format-specific properties ------------------------------------

    def _pdf_props(self, data: bytes) -> dict:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        info = reader.metadata or {}
        return _compact(
            {
                "title": info.title,
                "author": info.author,
                "subject": info.subject,
                "creator": info.creator,
                "producer": info.producer,
                "created": getattr(info, "creation_date", None),
                "modified": getattr(info, "modification_date", None),
            }
        )

    def _docx_props(self, data: bytes) -> dict:
        from docx import Document as DocxDocument

        cp = DocxDocument(io.BytesIO(data)).core_properties
        return _compact(
            {
                "title": cp.title,
                "author": cp.author,
                "subject": cp.subject,
                "keywords": cp.keywords,
                "language": cp.language,
                "category": cp.category,
                "last_modified_by": cp.last_modified_by,
                "revision": cp.revision,
                "created": cp.created,
                "modified": cp.modified,
            }
        )

    def _xlsx_props(self, data: bytes) -> dict:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True)
        p = wb.properties
        try:
            return _compact(
                {
                    "title": p.title,
                    "creator": p.creator,
                    "subject": p.subject,
                    "keywords": p.keywords,
                    "category": p.category,
                    "last_modified_by": p.lastModifiedBy,
                    "revision": p.revision,
                    "created": p.created,
                    "modified": p.modified,
                }
            )
        finally:
            wb.close()
